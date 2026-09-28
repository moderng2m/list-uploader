"""Generate the synthetic sample files (no real lead data, ever).

Run from the repo root:  uv run python backend/tests/fixtures/make_synthetic.py

Writes:
  backend/tests/fixtures/synthetic/*       parser test inputs
  frontend/public/List_Upload_Template.xlsx  blank template for the Upload page

The template's header row follows SPEC §8; the real TriNet template's column
order may differ.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent / "synthetic"
PUBLIC = ROOT / "frontend" / "public"
FIXED_TIME = datetime(2026, 9, 1, 12, 0, 0)

TEMPLATE_HEADERS = [
    "Company",
    "First name",
    "Last Name",
    "Email Address",
    "Lead Source - Most Recent",
    "SFDC Last Campaign ID",
    "Title",
    "Business Phone",
    "Mobile Phone",
    "Address 1",
    "City",
    "State or Province",
    "Zip or Postal Code",
    "Country",
    "Notes additional Information",
    "Zoom Individual ID",
    "Zoom Company ID",
    "SFDC Last Campaign Status",
    "NAICS Code",
    "Industry",
    "Website",
    "Number of Employees",
    "SFDC List Name",
    "Last Response Class",
    "SFDC Last Campaign Name",
]

INSTRUCTIONS = [
    ["How to use this template (synthetic stand-in)"],
    ["Fill in one lead per row on Sheet1. Required: Company, First name, Last Name,"],
    ["Email Address, SFDC Last Campaign ID, SFDC List Name."],
]

# Sheet1 grid as a user sees it. Numbers are stored as numbers in the XLSX
# (phone, ZIP, IDs, counts), exactly the case the parser must stringify.
# Row 2 holds only the stray #VALUE! in E2 that ships with the real template.
# Row 5 is blank.
_LEADS: list[list[object]] = [
    ["Acme Demo Co", "Ada", "Example", "ada@acme.example", "Marketing: Events",
     "701000000000001AAA", "VP Marketing", 5550100100, "+1 555-010-0199", "1 Demo Way",
     "Boston", "MA", 2134, "United States", "Met at booth 12", 1000000001, 2000000001,
     "Attended", 541511, "Software", "acme.example", 250, "Demo Conf 2026 - Booth",
     "", "Demo Conference 2026"],
    ["Globex Test Inc", "José", "Müller", "jose.muller@globex.example", "Events",
     "701000000000001AAA", "Director, Operations", 5550100101, "", "22 Sample St",
     "Chicago", "IL", 60601, "USA", "", "", "", "", 522110, "Financial Services",
     "https://globex.example", 1200, "Demo Conf 2026 - Booth", "", ""],
    None,  # blank row
    ["Initech Sample", "Zoë", "Placeholder", "zoe@initech.example", "", "701000000000001AAA",
     "CFO", "", "", "", "Austin", "TX", 73301, "US", "", "", "", "Registered", "", "", "",
     "10-50", "Demo Conf 2026 - Booth", "", ""],
    ["Umbrella Fake LLC", "Linus", "Sample", "linus@umbrella.example", "Marketing: Events",
     "701000000000002AAA", "Engineer", 5550100103, "", "", "", "", "", "Canada", "", "", "",
     "", "", "", "", 75, "Demo Webinar - Registrants", "", ""],
]  # fmt: skip


def template_grid() -> list[list[object]]:
    stray = [""] * len(TEMPLATE_HEADERS)
    stray[4] = "#VALUE!"
    grid: list[list[object]] = [list(TEMPLATE_HEADERS), stray]
    for lead in _LEADS:
        grid.append(lead if lead is not None else [""] * len(TEMPLATE_HEADERS))
    return grid


def _as_text(value: object) -> str:
    return "" if value is None else str(value)


def _stamp(wb: Workbook) -> None:
    wb.properties.creator = "list-uploader synthetic fixtures"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME


def write_template_xlsx(path: Path, *, with_data: bool) -> None:
    wb = Workbook()
    ins = wb.active
    assert ins is not None
    ins.title = "Instructions"
    for row in INSTRUCTIONS:
        ins.append(row)
    sheet = wb.create_sheet("Sheet1")
    rows = template_grid() if with_data else [TEMPLATE_HEADERS]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value != "":
                sheet.cell(row=r, column=c, value=value)
    for c in range(1, len(TEMPLATE_HEADERS) + 1):
        sheet.cell(row=1, column=c).font = Font(bold=True)
    _stamp(wb)
    wb.save(path)


def write_csv(path: Path, rows: list[list[object]], *, encoding: str, delimiter: str = ",") -> None:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    for row in rows:
        writer.writerow([_as_text(v) for v in row])
    path.write_bytes(buf.getvalue().encode(encoding))


def write_vendor_xlsx(path: Path) -> None:
    """A messy vendor export: odd headers, duplicates, a blank header, dates, floats."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Export"
    # Two blank rows before the header.
    headers = ["E-mail", "Org", "Job Position", "First", "Last", "Phone", "Phone", None,
               "Scan Date", "Zip", "Score", "Opt In"]  # fmt: skip
    data = [
        ["ada@acme.example", "Acme Demo Co", "VP Marketing", "Ada", "Example", 5550100100.0,
         "555-010-0198", "booth A", datetime(2026, 9, 15), 2134, 12.5, True],
        ["grace@globex.example", "Globex Test Inc", "Director", "Grace", "Sample", None,
         None, None, datetime(2026, 9, 15, 14, 30), 501, 7, False],
        ["test@test.com", "asdf", "n/a", "Test", "Test", "", "", "", None, "", None, None],
    ]  # fmt: skip
    for c, value in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=value)
    for r, row in enumerate(data, start=4):
        for c, value in enumerate(row, start=1):
            if value not in (None, ""):
                ws.cell(row=r, column=c, value=value)
    ws.cell(row=4, column=9).number_format = "m/d/yyyy"
    # Formatting far below and to the right, with no values: must be ignored.
    ws.cell(row=60, column=3).font = Font(bold=True)
    ws.cell(row=5, column=20).font = Font(italic=True)
    _stamp(wb)
    wb.save(path)


def cp1252_rows() -> list[list[object]]:
    return [
        ["Company", "First name", "Last Name", "Email Address", "Notes additional Information"],
        ["Société Générale Démo", "José", "Müller", "jose@societe.example",
         "Prefers “email” – no calls"],
        ["Café Zoë", "Zoë", "Brontë", "zoe@cafe.example", "Price in € noted"],
    ]  # fmt: skip


def semicolon_rows() -> list[list[object]]:
    return [
        ["Company", "First name", "Last Name", "Email Address", "Number of Employees", "Notes"],
        ["Acme, Demo GmbH", "Anna", "Beispiel", "anna@example.de", "1.200", "Score 4,5"],
        ["Beispiel AG", "Max", "Muster", "max@example.de", "80", ""],
    ]  # fmt: skip


def analysis_rows() -> list[list[object]]:
    """P3 acceptance file: campaign ID variants, statuses, lead sources, junk, duplicates."""
    header = ["Company", "First name", "Last Name", "Email Address", "SFDC Last Campaign ID",
              "SFDC Last Campaign Status", "Lead Source - Most Recent", "Business Phone",
              "Zip or Postal Code", "Country"]  # fmt: skip
    return [
        header,
        # 2: 15-char ID, blank status -> default, "Events" -> "Marketing: Events"
        ["Acme Demo Co", "Ada", "Example", "ada@acme.example", "701000000000001", "", "Events",
         "555-010-0100", "2134", "USA"],
        # 3: bad checksum (retyped suffix)
        ["Globex Test Inc", "Grace", "Sample", "grace@globex.example", "701000000000001AAB",
         "Attended", "Marketing: Events", "", "", ""],
        # 4: valid ID that Salesforce doesn't have
        ["Initech Sample", "Alan", "Placeholder", "alan@initech.example", "701000000000009AAA",
         "Registered", "Marketing: Events", "", "", ""],
        # 5: duplicate of row 2 (same email, same campaign in 18-char form)
        ["Acme Demo Co", "Ada", "Example", "ADA@acme.example", "701000000000001AAA", "Attended",
         "Marketing: Events", "", "", ""],
        # 6: junk
        ["asdf", "Test", "Test", "test@umbrella.example", "701000000000001AAA", "Attended",
         "Marketing: Events", "", "", ""],
        # 7: webinar campaign, misspelled status, lead source that needs the AI step
        ["Hooli Example", "Linus", "Sample", "linus@hooli.example", "701000000000002AAA",
         "Atended", "Webcast", "", "", ""],
        # 8: blank company (pending when enrichment is on)
        ["", "Kay", "Sample", "kay@pied.example", "701000000000001AAA", "Registered",
         "Marketing: Events", "", "", ""],
        # 9: inactive campaign, role inbox, bad phone
        ["Vandelay Demo", "Art", "Sample", "info@vandelay.example", "701000000000003AAA",
         "Attended", "Marketing: Events", "12345", "", ""],
    ]  # fmt: skip


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_template_xlsx(OUT / "template_filled.xlsx", with_data=True)
    write_csv(OUT / "template_filled_utf8_bom.csv", template_grid(), encoding="utf-8-sig")
    write_csv(OUT / "windows_1252.csv", cp1252_rows(), encoding="cp1252")
    write_csv(OUT / "semicolon.csv", semicolon_rows(), encoding="utf-8", delimiter=";")
    write_vendor_xlsx(OUT / "vendor_export.xlsx")
    write_csv(OUT / "analysis_demo.csv", analysis_rows(), encoding="utf-8")
    PUBLIC.mkdir(parents=True, exist_ok=True)
    write_template_xlsx(PUBLIC / "List_Upload_Template.xlsx", with_data=False)
    print(f"wrote fixtures to {OUT} and template to {PUBLIC}")


if __name__ == "__main__":
    main()
