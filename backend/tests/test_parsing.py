from __future__ import annotations

import io
from datetime import datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook

from shared.parsing import ParseError, ParseResult, cell_to_str, parse_file, sniff_delimiter
from tests.fixtures.make_synthetic import TEMPLATE_HEADERS

FIXTURES = Path(__file__).parent / "fixtures" / "synthetic"


def _parse(name: str) -> ParseResult:
    return parse_file((FIXTURES / name).read_bytes(), name)


def _xlsx(rows: list[list[object]], sheets: tuple[str, ...] = ("Sheet1",)) -> bytes:
    wb = Workbook()
    first = wb.active
    assert first is not None
    first.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    target = wb["Sheet1"] if "Sheet1" in sheets else first
    for row in rows:
        target.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(result: ParseResult) -> list[tuple[int, dict[str, str]]]:
    return [(r.row_id, r.values) for r in result.rows]


class TestTemplate:
    def test_template_xlsx(self) -> None:
        result = _parse("template_filled.xlsx")
        assert result.file_type == "xlsx"
        assert result.sheet_name == "Sheet1"  # Instructions sheet comes first and is skipped
        assert result.headers == TEMPLATE_HEADERS
        assert [r.row_id for r in result.rows] == [3, 4, 6, 7]  # row 2 (stray) and 5 dropped

    def test_stray_value_error_in_e2_is_blanked_and_warned(self) -> None:
        result = _parse("template_filled.xlsx")
        assert [(w.code, w.row_id, w.column) for w in result.warnings] == [
            ("EXCEL_ERROR_VALUE", 2, "Lead Source - Most Recent")
        ]
        assert "E2" in result.warnings[0].message
        assert all("#VALUE!" not in v for r in result.rows for v in r.values.values())

    def test_numbers_stored_as_numbers_become_plain_strings(self) -> None:
        first = _parse("template_filled.xlsx").rows[0].values
        assert first["Business Phone"] == "5550100100"
        assert first["Zip or Postal Code"] == "2134"  # leading zero restored after mapping (§7.1)
        assert first["Zoom Individual ID"] == "1000000001"
        assert first["NAICS Code"] == "541511"
        assert first["Number of Employees"] == "250"

    def test_same_data_as_csv_and_xlsx_parses_identically(self) -> None:
        xlsx = _parse("template_filled.xlsx")
        csv = _parse("template_filled_utf8_bom.csv")
        assert csv.encoding == "utf-8-sig"
        assert csv.headers == xlsx.headers
        assert _rows(csv) == _rows(xlsx)
        assert [r.issues for r in csv.rows] == [r.issues for r in xlsx.rows]
        assert csv.warnings == xlsx.warnings

    def test_unicode_survives(self) -> None:
        second = _parse("template_filled_utf8_bom.csv").rows[1].values
        assert (second["First name"], second["Last Name"]) == ("José", "Müller")


class TestCsv:
    def test_windows_1252(self) -> None:
        result = _parse("windows_1252.csv")
        assert result.encoding == "windows-1252"
        first = result.rows[0].values
        assert first["Company"] == "Société Générale Démo"
        assert first["Notes additional Information"] == "Prefers “email” – no calls"
        assert result.rows[1].values["Notes additional Information"] == "Price in € noted"

    def test_semicolon(self) -> None:
        result = _parse("semicolon.csv")
        assert result.delimiter == ";"
        assert result.rows[0].values["Company"] == "Acme, Demo GmbH"
        assert result.rows[0].values["Notes"] == "Score 4,5"

    def test_tab_delimited_utf16(self) -> None:
        # Excel's "Unicode Text" export.
        data = "Company\tEmail Address\r\nAcme\tada@example.com\r\n".encode("utf-16")
        result = parse_file(data, "export.csv")
        assert (result.encoding, result.delimiter) == ("utf-16", "\t")
        assert result.rows[0].values == {"Company": "Acme", "Email Address": "ada@example.com"}

    def test_values_are_never_type_inferred(self) -> None:
        data = b"Zip,Phone,Flag,Id\n02134,+15550100100,NA,0012\n"
        values = parse_file(data, "f.csv").rows[0].values
        assert values == {"Zip": "02134", "Phone": "+15550100100", "Flag": "NA", "Id": "0012"}

    def test_quoted_newlines_and_row_numbers(self) -> None:
        data = b'Company,Notes\n"Acme","line one\nline two"\n\nGlobex,x\n'
        result = parse_file(data, "f.csv")
        assert [(r.row_id, r.values["Notes"]) for r in result.rows] == [
            (2, "line one\nline two"),
            (4, "x"),
        ]

    def test_ragged_rows_pad_and_extra_cells_get_named_column(self) -> None:
        data = b"A,B\n1\n2,3,4\n"
        result = parse_file(data, "f.csv")
        assert result.headers == ["A", "B", "Column C"]
        assert _rows(result) == [(2, {"A": "1", "B": "", "Column C": ""}),
                                 (3, {"A": "2", "B": "3", "Column C": "4"})]  # fmt: skip

    @pytest.mark.parametrize(
        ("line", "expected"),
        [("a,b,c", ","), ("a;b;c", ";"), ("a\tb\tc", "\t"), ('"x;y",b,c', ","), ("single", ",")],
    )
    def test_sniff_delimiter(self, line: str, expected: str) -> None:
        assert sniff_delimiter(f"{line}\n{line}\n") == expected


class TestXlsx:
    def test_vendor_export(self) -> None:
        result = _parse("vendor_export.xlsx")
        assert result.sheet_name == "Export"
        assert result.headers == [
            "E-mail", "Org", "Job Position", "First", "Last", "Phone", "Phone (2)",
            "Column H", "Scan Date", "Zip", "Score", "Opt In",
        ]  # fmt: skip
        assert {w.code for w in result.warnings} == {"DUPLICATE_HEADER", "BLANK_HEADER"}
        assert [r.row_id for r in result.rows] == [4, 5, 6]  # header on row 3; formatting ignored
        ada, grace = result.rows[0].values, result.rows[1].values
        assert ada["Phone"] == "5550100100"
        assert ada["Scan Date"] == "2026-09-15"
        assert grace["Scan Date"] == "2026-09-15T14:30:00"
        assert ada["Score"] == "12.5"
        assert ada["Opt In"] == "TRUE"
        assert grace["Zip"] == "501"

    def test_sheet1_preferred_over_first_sheet(self) -> None:
        data = _xlsx([["Company"], ["Acme"]], sheets=("Other", "Sheet1"))
        assert parse_file(data, "f.xlsx").sheet_name == "Sheet1"

    def test_only_instructions_sheet_is_empty(self) -> None:
        wb = Workbook()
        assert wb.active is not None
        wb.active.title = "Instructions"
        wb.active.append(["Read me"])
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ParseError) as err:
            parse_file(buf.getvalue(), "f.xlsx")
        assert err.value.code == "EMPTY_FILE"

    def test_duplicate_headers_are_case_insensitive(self) -> None:
        result = parse_file(_xlsx([["Email", "email", "EMAIL"], ["a", "b", "c"]]), "f.xlsx")
        assert result.headers == ["Email", "email (2)", "EMAIL (3)"]

    def test_headers_trimmed_and_empty_columns_dropped(self) -> None:
        result = parse_file(
            _xlsx([["  Company ", None, "Email"], ["Acme", None, "a@b.co"]]), "f.xlsx"
        )
        assert result.headers == ["Company", "Email"]

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            (5551234567.0, "5551234567"),
            (0.1 + 0.2, "0.3"),
            (1e20, "1e+20"),
            (-3.0, "-3"),
            (datetime(2026, 1, 2), "2026-01-02"),
            (time(9, 30), "09:30:00"),
            (None, ""),
            (False, "FALSE"),
        ],
    )
    def test_cell_to_str(self, value: object, expected: str) -> None:
        assert cell_to_str(value) == expected


class TestRejections:
    def test_oversize_file(self) -> None:
        with pytest.raises(ParseError) as err:
            parse_file(b"a" * 101, "big.csv", max_bytes=100)
        assert err.value.code == "FILE_TOO_LARGE"

    def test_default_limit_is_10_mb(self) -> None:
        data = b"Company\n" + b"x" * (10 * 1024 * 1024)
        with pytest.raises(ParseError) as err:
            parse_file(data, "big.csv")
        assert err.value.code == "FILE_TOO_LARGE"
        assert "10 MB" in err.value.message

    def test_too_many_rows(self) -> None:
        data = ("Company\n" + "Acme\n" * 6).encode()
        with pytest.raises(ParseError) as err:
            parse_file(data, "f.csv", max_rows=5)
        assert err.value.code == "TOO_MANY_ROWS"
        assert parse_file(("Company\n" + "Acme\n" * 5).encode(), "f.csv", max_rows=5)

    @pytest.mark.parametrize(
        ("data", "code"),
        [
            (b"", "EMPTY_FILE"),
            (b"\n\n , ,\n", "EMPTY_FILE"),
            (b"Company,Email\n", "NO_DATA_ROWS"),
            (b"Company\n#VALUE!\n", "NO_DATA_ROWS"),
        ],
    )
    def test_empty_inputs(self, data: bytes, code: str) -> None:
        with pytest.raises(ParseError) as err:
            parse_file(data, "f.csv")
        assert err.value.code == code

    def test_wrong_extension(self) -> None:
        with pytest.raises(ParseError) as err:
            parse_file(b"x", "leads.pdf")
        assert err.value.code == "WRONG_FILE_TYPE"
        assert "This file is a .pdf" in err.value.message

    def test_password_protected_xlsx(self) -> None:
        ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
        with pytest.raises(ParseError) as err:
            parse_file(ole, "locked.xlsx")
        assert err.value.code == "XLSX_PROTECTED"

    def test_xlsx_named_csv(self) -> None:
        with pytest.raises(ParseError) as err:
            parse_file(_xlsx([["Company"], ["Acme"]]), "leads.csv")
        assert err.value.code == "CSV_IS_EXCEL"

    def test_corrupt_xlsx(self) -> None:
        for data in (b"PK\x03\x04garbage", b"just text"):
            with pytest.raises(ParseError) as err:
                parse_file(data, "leads.xlsx")
            assert err.value.code == "XLSX_UNREADABLE"
