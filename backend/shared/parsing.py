"""CSV and XLSX parsing (SPEC §6.1, §7.1).

Every value comes out as a string exactly as the user sees it: no type
inference, integers stored as floats lose their `.0`, Excel error values
become blank with a warning. The same data saved as CSV and as XLSX parses
to the same result.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from shared import config_defaults, messages

EXCEL_ERRORS = frozenset(
    {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!"}
)
IGNORED_SHEETS = frozenset({"instructions"})
PREFERRED_SHEET = "sheet1"
CSV_DELIMITERS = (",", ";", "\t")
MAX_COLUMNS = 200
MAX_ROW_CHARS = 100_000  # keeps each Rows item well under DynamoDB's 400 KB

_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # encrypted .xlsx or legacy .xls


class ParseError(Exception):
    """The file can't be used. `message` is shown to the user as-is."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ParseIssue:
    code: str
    severity: str
    message: str
    row_id: int | None = None
    column: str | None = None

    def as_dict(self) -> dict[str, Any]:
        out: dict[str, Any] = {
            "code": self.code,
            "severity": self.severity,
            "message": self.message,
        }
        if self.row_id is not None:
            out["row_id"] = self.row_id
        if self.column is not None:
            out["column"] = self.column
        return out


@dataclass(frozen=True)
class ParsedRow:
    row_id: int  # 1-based row number in the source file
    values: dict[str, str]
    issues: tuple[ParseIssue, ...] = ()


@dataclass
class ParseResult:
    file_type: str  # "csv" | "xlsx"
    headers: list[str]
    rows: list[ParsedRow]
    warnings: list[ParseIssue] = field(default_factory=list)  # file-level
    sheet_name: str | None = None
    encoding: str | None = None
    delimiter: str | None = None

    def row_issue_count(self) -> int:
        return sum(len(r.issues) for r in self.rows)


# --- public API ----------------------------------------------------------------


def parse_file(
    data: bytes,
    filename: str,
    *,
    max_bytes: int = config_defaults.LIMITS["max_file_bytes"],
    max_rows: int = config_defaults.LIMITS["max_rows"],
) -> ParseResult:
    ext = file_extension(filename)
    if ext not in (".csv", ".xlsx"):
        raise ParseError("WRONG_FILE_TYPE", messages.WRONG_FILE_TYPE.format(ext=ext.lstrip(".")))
    if len(data) > max_bytes:
        raise ParseError(
            "FILE_TOO_LARGE", messages.FILE_TOO_LARGE.format(limit_mb=max_bytes // (1024 * 1024))
        )
    if not data.strip():
        raise ParseError("EMPTY_FILE", messages.EMPTY_FILE)

    # Trust the bytes, not the extension.
    if data.startswith(_OLE_MAGIC):
        raise ParseError("XLSX_PROTECTED", messages.XLSX_PROTECTED)
    if data.startswith(_ZIP_MAGIC):
        if ext == ".csv":
            raise ParseError("CSV_IS_EXCEL", messages.CSV_IS_EXCEL)
        grid, sheet_name = _read_xlsx(data)
        return _build("xlsx", grid, max_rows, sheet_name=sheet_name)
    if ext == ".xlsx":
        raise ParseError("XLSX_UNREADABLE", messages.XLSX_UNREADABLE)

    text, encoding = decode_text(data)
    delimiter = sniff_delimiter(text)
    grid = _read_csv(text, delimiter)
    return _build("csv", grid, max_rows, encoding=encoding, delimiter=delimiter)


def file_extension(filename: str) -> str:
    name = filename.strip().lower()
    return name[name.rfind(".") :] if "." in name else ""


# --- CSV -------------------------------------------------------------------------


def decode_text(data: bytes) -> tuple[str, str]:
    """UTF-8 (with or without BOM), UTF-16 with BOM, else Windows-1252."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252"), "windows-1252"
    except UnicodeDecodeError:
        # cp1252 leaves 5 bytes undefined; latin-1 maps every byte.
        return data.decode("latin-1"), "latin-1"


def sniff_delimiter(text: str, sample_lines: int = 50) -> str:
    """Pick the delimiter that gives the most columns, consistently, over the first lines."""
    sample = "\n".join(line for line in text.splitlines()[:sample_lines] if line.strip())
    best, best_score = ",", (0, 0)
    for delim in CSV_DELIMITERS:
        widths = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=delim) if any(r)]
        if not widths:
            continue
        header_width = widths[0]
        consistent = sum(1 for w in widths if w == header_width)
        score = (header_width, consistent) if header_width > 1 else (0, 0)
        if score > best_score:
            best, best_score = delim, score
    return best


def _read_csv(text: str, delimiter: str) -> Iterator[tuple[int, list[str]]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    yield from enumerate(reader, start=1)


# --- XLSX ------------------------------------------------------------------------


def _read_xlsx(data: bytes) -> tuple[Iterator[tuple[int, list[str]]], str]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ParseError("XLSX_UNREADABLE", messages.XLSX_UNREADABLE) from exc

    names = wb.sheetnames
    by_lower = {n.lower(): n for n in names}
    candidates = [n for n in names if n.lower() not in IGNORED_SHEETS]
    sheet_name = by_lower.get(PREFERRED_SHEET) or (candidates[0] if candidates else None)
    if sheet_name is None:
        raise ParseError("EMPTY_FILE", messages.EMPTY_FILE)
    ws = wb[sheet_name]

    def rows() -> Iterator[tuple[int, list[str]]]:
        try:
            for index, values in enumerate(ws.iter_rows(values_only=True), start=1):
                yield index, [cell_to_str(v) for v in values]
        finally:
            wb.close()

    return rows(), sheet_name


def cell_to_str(value: Any) -> str:
    """Excel cell value -> the string a user would see in a CSV export."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return format(value, ".15g")
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


# --- shared row assembly -----------------------------------------------------------


def _build(
    file_type: str,
    grid: Iterable[tuple[int, list[str]]],
    max_rows: int,
    *,
    sheet_name: str | None = None,
    encoding: str | None = None,
    delimiter: str | None = None,
) -> ParseResult:
    grid_iter = iter(grid)

    # Header = first non-empty row.
    header_row: list[str] | None = None
    for _, cells in grid_iter:
        if any(c.strip() for c in cells):
            header_row = [c.strip() for c in cells]
            break
    if header_row is None:
        raise ParseError("EMPTY_FILE", messages.EMPTY_FILE)

    raw_rows: list[tuple[int, list[str]]] = []
    for index, cells in grid_iter:
        if any(c.strip() for c in cells):
            raw_rows.append((index, cells))
            if len(raw_rows) > max_rows:
                raise ParseError("TOO_MANY_ROWS", messages.TOO_MANY_ROWS.format(limit=max_rows))

    width = max([len(header_row)] + [len(c) for _, c in raw_rows])
    headers, keep, warnings = _resolve_headers(header_row, raw_rows, width)
    if len(headers) > MAX_COLUMNS:
        raise ParseError("TOO_MANY_COLUMNS", messages.TOO_MANY_COLUMNS.format(limit=MAX_COLUMNS))

    rows: list[ParsedRow] = []
    for index, cells in raw_rows:
        values: dict[str, str] = {}
        issues: list[ParseIssue] = []
        for col, header in zip(keep, headers, strict=True):
            value = cells[col] if col < len(cells) else ""
            if value.strip() in EXCEL_ERRORS:
                issues.append(
                    ParseIssue(
                        code="EXCEL_ERROR_VALUE",
                        severity="warning",
                        message=messages.EXCEL_ERROR_VALUE.format(
                            value=value.strip(), cell=f"{get_column_letter(col + 1)}{index}"
                        ),
                        row_id=index,
                        column=header,
                    )
                )
                value = ""
            values[header] = value
        if sum(len(v) for v in values.values()) > MAX_ROW_CHARS:
            raise ParseError("ROW_TOO_LARGE", messages.ROW_TOO_LARGE.format(row=index))
        if any(v.strip() for v in values.values()):
            rows.append(ParsedRow(row_id=index, values=values, issues=tuple(issues)))
        else:
            # Row held nothing but error values: drop it, keep the warnings.
            warnings.extend(issues)

    if not rows:
        raise ParseError("NO_DATA_ROWS", messages.NO_DATA_ROWS)
    return ParseResult(
        file_type=file_type,
        headers=headers,
        rows=rows,
        warnings=warnings,
        sheet_name=sheet_name,
        encoding=encoding,
        delimiter=delimiter,
    )


def _resolve_headers(
    header_row: list[str], raw_rows: list[tuple[int, list[str]]], width: int
) -> tuple[list[str], list[int], list[ParseIssue]]:
    """Trim, name blank-but-used columns, drop blank-and-empty columns, de-duplicate."""
    headers: list[str] = []
    keep: list[int] = []
    warnings: list[ParseIssue] = []
    seen: dict[str, int] = {}
    for col in range(width):
        name = header_row[col] if col < len(header_row) else ""
        letter = get_column_letter(col + 1)
        if not name:
            has_data = any(col < len(c) and c[col].strip() for _, c in raw_rows)
            if not has_data:
                continue
            name = f"Column {letter}"
            warnings.append(
                ParseIssue(
                    "BLANK_HEADER",
                    "warning",
                    messages.BLANK_HEADER.format(letter=letter),
                    None,
                    name,
                )
            )
        key = name.casefold()
        if key in seen:
            seen[key] += 1
            original = name
            name = f"{name} ({seen[key]})"
            warnings.append(
                ParseIssue(
                    "DUPLICATE_HEADER",
                    "warning",
                    messages.DUPLICATE_HEADER.format(header=original, renamed=name),
                    None,
                    name,
                )
            )
        else:
            seen[key] = 1
        headers.append(name)
        keep.append(col)
    return headers, keep, warnings
